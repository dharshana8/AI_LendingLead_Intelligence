from fastapi import FastAPI, Depends, HTTPException, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import csv
import io
import os
from groq import Groq
from dotenv import load_dotenv

load_dotenv()

from database import leads_col, users_col, get_next_sequence_value
from schemas import CustomerCreate, CustomerUpdate, CustomerResponse, AnalyticsResponse
from crud import create_lead, get_leads, get_lead, delete_lead, update_lead
from analytics import get_analytics
from predict import _load_artefacts
from sample_data import SAMPLE_CUSTOMERS


app = FastAPI(
    title="AI Lending Lead Intelligence API",
    description="IDBI Bank Hackathon — AI-powered lead scoring and loan recommendation backend.",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup_event():
    _load_artefacts()
    await _seed_users()
    await _seed_customers()


@app.get("/health")
def health():
    return {"status": "OK"}


# ── Auth ──────────────────────────────────────────────────────────────────────
DEFAULT_USERS = [
    {"employee_id": "RM001",  "password": "password123", "name": "Ankit Sharma", "role": "rm",    "branch": "Mumbai Main",  "email": "ankit.sharma@idbi.co.in", "phone": "+91 98765 43210", "avatar": "AS", "joined": "Jan 2022", "performance": 92},
    {"employee_id": "BM001",  "password": "password123", "name": "Priya Mehta",  "role": "bm",    "branch": "Delhi Central", "email": "priya.mehta@idbi.co.in",  "phone": "+91 98765 43211", "avatar": "PM", "joined": "Mar 2019", "performance": 88},
    {"employee_id": "ADM001", "password": "admin123",    "name": "Rajiv Nair",   "role": "admin", "branch": "HQ Mumbai",     "email": "rajiv.nair@idbi.co.in",   "phone": "+91 98765 43212", "avatar": "RN", "joined": "Jun 2015", "performance": 96},
]


async def _seed_users():
    for u in DEFAULT_USERS:
        exists = await users_col.find_one({"employee_id": u["employee_id"]})
        if not exists:
            uid = await get_next_sequence_value("user_id")
            await users_col.insert_one({**u, "id": uid})


async def _seed_customers() -> list[dict]:
    """Seed sample customers only when the leads collection is empty."""
    if await leads_col.count_documents({}) > 0:
        return await get_leads()
    return await _insert_sample_customers()


async def _insert_sample_customers() -> list[dict]:
    results = []
    for customer in SAMPLE_CUSTOMERS:
        payload = CustomerCreate(**customer)
        res = await create_lead(payload)
        results.append(res)
    return results


def _user_dict(u: dict):
    return {"id": u.get("id"), "employeeId": u.get("employee_id"), "name": u.get("name"), "role": u.get("role"),
            "branch": u.get("branch"), "email": u.get("email"), "phone": u.get("phone"),
            "avatar": u.get("avatar"), "joined": u.get("joined"), "performance": u.get("performance")}


class LoginRequest(BaseModel):
    employeeId: str
    password: str


class RegisterRequest(BaseModel):
    employeeId: str
    password: str
    name: str
    role: str = "rm"
    branch: str = "Mumbai Main"
    email: str = ""
    phone: str = ""


@app.post("/login")
async def login(payload: LoginRequest):
    user = await users_col.find_one({
        "$or": [
            {"employee_id": payload.employeeId, "password": payload.password},
            {"email": payload.employeeId, "password": payload.password},
        ]
    })
    if not user:
        raise HTTPException(status_code=401, detail="Invalid Employee ID / Email or Password")
    return _user_dict(user)


@app.post("/register", status_code=201)
async def register(payload: RegisterRequest):
    existing = await users_col.find_one({"employee_id": payload.employeeId})
    if existing:
        raise HTTPException(status_code=409, detail="Employee ID already exists")
    uid = await get_next_sequence_value("user_id")
    new_user = {
        "id": uid,
        "employee_id": payload.employeeId,
        "password": payload.password,
        "name": payload.name,
        "role": payload.role,
        "branch": payload.branch,
        "email": payload.email,
        "phone": payload.phone,
        "avatar": "".join(w[0] for w in payload.name.split()[:2]).upper(),
        "joined": "2025",
        "performance": 75,
    }
    await users_col.insert_one(new_user)
    return _user_dict(new_user)


# ── Customers ─────────────────────────────────────────────────────────────────
@app.post("/customers", response_model=CustomerResponse, status_code=201)
async def create_customer(payload: CustomerCreate):
    return await create_lead(payload)


@app.get("/customers", response_model=list[CustomerResponse])
async def list_customers(
    priority: str | None = Query(default=None, pattern="^(High|Medium|Low)$"),
):
    return await get_leads(priority)


@app.get("/customers/{customer_id}", response_model=CustomerResponse)
async def get_customer(customer_id: int):
    lead = await get_lead(customer_id)
    if not lead:
        raise HTTPException(status_code=404, detail="Customer not found.")
    return lead


@app.delete("/customers/{customer_id}", status_code=204)
async def remove_customer(customer_id: int):
    if not await delete_lead(customer_id):
        raise HTTPException(status_code=404, detail="Customer not found.")


@app.put("/customers/{customer_id}", response_model=CustomerResponse)
async def update_customer(customer_id: int, payload: CustomerUpdate):
    lead = await update_lead(customer_id, payload)
    if not lead:
        raise HTTPException(status_code=404, detail="Customer not found.")
    return lead


@app.post("/load-sample", response_model=list[CustomerResponse])
async def load_sample(force: bool = Query(default=False)):
    """Load sample customers. Skips if data already exists unless force=true."""
    if not force and await leads_col.count_documents({}) > 0:
        return await get_leads()
    if force:
        await leads_col.delete_many({})
    return await _insert_sample_customers()


@app.get("/export")
async def export_csv():
    leads = await get_leads()
    output = io.StringIO()
    fields = [
        "customer_id", "name", "age", "occupation", "cibil_score",
        "monthly_credit_1", "monthly_credit_2", "monthly_credit_3",
        "monthly_credit_4", "monthly_credit_5", "monthly_credit_6",
        "emi_debits", "cc_spend", "credit_limit", "account_balance",
        "existing_loan_count", "years_of_experience", "loan_page_visits",
        "income", "salary_regularity", "emi_burden", "savings_ratio",
        "ai_score", "conversion_probability", "priority", "recommended_loan",
        "top_signal", "status", "assigned_to", "last_contact", "created_at",
    ]
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for lead in leads:
        row = {f: lead.get(f, "") for f in fields}
        writer.writerow(row)
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=idbi_leads_export.csv"},
    )


@app.get("/analytics", response_model=AnalyticsResponse)
async def analytics():
    return await get_analytics()


@app.post("/import-csv")
async def import_csv(file: UploadFile = File(...)):
    """Import customers from CSV or Excel file. Maps common column name variants."""
    filename = file.filename.lower()
    content = await file.read()

    rows = []
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        except ImportError:
            raise HTTPException(status_code=400, detail="openpyxl not installed. Use CSV format.")
    else:
        raise HTTPException(status_code=400, detail="Only .csv, .xlsx, .xls files supported")

    if not rows:
        raise HTTPException(status_code=400, detail="File is empty or has no data rows")

    # Column name aliases — maps various header names → our field names
    ALIASES = {
        "name": ["name", "customer_name", "full_name", "customer name"],
        "age": ["age"],
        "occupation": ["occupation", "job", "job_title", "profession"],
        "cibil_score": ["cibil_score", "cibil", "credit_score"],
        "monthly_credit_1": ["monthly_credit_1", "month1", "m1", "salary_m1"],
        "monthly_credit_2": ["monthly_credit_2", "month2", "m2", "salary_m2"],
        "monthly_credit_3": ["monthly_credit_3", "month3", "m3", "salary_m3"],
        "monthly_credit_4": ["monthly_credit_4", "month4", "m4", "salary_m4"],
        "monthly_credit_5": ["monthly_credit_5", "month5", "m5", "salary_m5"],
        "monthly_credit_6": ["monthly_credit_6", "month6", "m6", "salary_m6"],
        "emi_debits": ["emi_debits", "emi", "monthly_emi"],
        "cc_spend": ["cc_spend", "credit_card_spend", "cc"],
        "credit_limit": ["credit_limit", "limit"],
        "loan_page_visits": ["loan_page_visits", "page_visits", "visits"],
        "existing_loan_count": ["existing_loan_count", "existing_loans", "loan_count"],
        "years_of_experience": ["years_of_experience", "experience", "exp", "years_exp"],
        "account_balance": ["account_balance", "balance", "bank_balance"],
    }

    def resolve(row: dict, field: str):
        row_lower = {k.lower().strip(): v for k, v in row.items()}
        for alias in ALIASES.get(field, [field]):
            if alias in row_lower:
                return row_lower[alias]
        return ""

    def safe_int(v, default=0):
        try: return int(float(str(v).replace(",", "").strip())) if v else default
        except: return default

    def safe_float(v, default=0.0):
        try: return float(str(v).replace(",", "").strip()) if v else default
        except: return default

    imported, skipped = [], []
    for i, row in enumerate(rows):
        name = resolve(row, "name").strip()
        if not name:
            skipped.append(f"Row {i+2}: missing name")
            continue
        cibil = safe_int(resolve(row, "cibil_score"), 650)
        if cibil < 300 or cibil > 900:
            cibil = 650
        credits = [safe_float(resolve(row, f"monthly_credit_{n}")) for n in range(1, 7)]
        # If no monthly credits provided, try to use income/salary column as credit_1
        if all(c == 0 for c in credits):
            row_lower = {k.lower().strip(): v for k, v in row.items()}
            for col in ["income", "monthly_income", "salary", "monthly_salary"]:
                if col in row_lower and safe_float(row_lower[col]) > 0:
                    credits[0] = safe_float(row_lower[col])
                    credits[1] = credits[0]
                    credits[2] = credits[0]
                    break
        try:
            payload = CustomerCreate(
                name=name,
                age=safe_int(resolve(row, "age"), 35),
                occupation=resolve(row, "occupation") or "Professional",
                cibil_score=cibil,
                monthly_credit_1=credits[0], monthly_credit_2=credits[1],
                monthly_credit_3=credits[2], monthly_credit_4=credits[3],
                monthly_credit_5=credits[4], monthly_credit_6=credits[5],
                emi_debits=safe_float(resolve(row, "emi_debits")),
                cc_spend=safe_float(resolve(row, "cc_spend")),
                credit_limit=safe_float(resolve(row, "credit_limit")) or 100000.0,
                loan_page_visits=safe_int(resolve(row, "loan_page_visits")),
                existing_loan_count=safe_int(resolve(row, "existing_loan_count")),
                years_of_experience=safe_int(resolve(row, "years_of_experience")),
                account_balance=safe_float(resolve(row, "account_balance")),
            )
            result = await create_lead(payload)
            imported.append(result)
        except Exception as e:
            skipped.append(f"Row {i+2} ({name}): {str(e)}")

    return {"imported": len(imported), "skipped": len(skipped), "skipped_details": skipped[:10]}


# ── Groq AI Chat ──────────────────────────────────────────────────────────────
class ChatRequest(BaseModel):
    message: str
    history: list[dict] = []


@app.post("/chat")
async def chat(payload: ChatRequest):
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="GROQ_API_KEY not configured")

    # Build live customer context from DB
    leads = await get_leads()
    top5 = leads[:5]
    lead_summary = "\n".join(
        f"- {l.get('name')} | {l.get('occupation')} | AI Score: {l.get('ai_score', 0.0):.0f} | "
        f"Priority: {l.get('priority')} | Loan: {l.get('recommended_loan')} | "
        f"Income: Rs{l.get('income', 0.0):,.0f} | CIBIL: {l.get('cibil_score')} | "
        f"Conversion: {l.get('conversion_probability', 0.0)*100:.0f}%"
        for l in top5
    )
    total = len(leads)
    high = sum(1 for l in leads if l.get("priority") == "High")
    avg_score = round(sum(l.get("ai_score", 0.0) for l in leads) / total, 1) if total else 0

    system_prompt = f"""You are an AI assistant for IDBI Bank's Lending Lead Intelligence platform.
You help Relationship Managers identify, prioritize, and convert high-value loan leads.

Current portfolio snapshot:
- Total customers: {total}
- High priority leads: {high}
- Average AI score: {avg_score}

Top leads right now:
{lead_summary}

Guidelines:
- Be concise, actionable, and specific to IDBI Bank context
- Use Indian currency (Rs) and Indian banking terminology
- Reference actual customer data above when relevant
- Suggest specific outreach strategies and loan products
- Keep responses focused on lending lead intelligence"""

    messages = [{"role": "system", "content": system_prompt}]
    for h in payload.history[-10:]:  # last 10 messages for context
        messages.append({"role": h["role"], "content": h["content"]})
    messages.append({"role": "user", "content": payload.message})

    try:
        client = Groq(api_key=api_key)
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=messages,
            max_tokens=1024,
            temperature=0.7,
        )
        return {"reply": response.choices[0].message.content}
    except Exception as e:
        raise HTTPException(status_code=503, detail=f"AI Assistant error: {str(e)}")

